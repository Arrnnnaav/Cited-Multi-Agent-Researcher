"""Convert the project deep-dive Markdown into a formatted .docx.

Handles: # / ## / ### headings, pipe tables, fenced code blocks, bullet lists,
blockquotes, horizontal rules, and **bold** / `code` inline spans. Tailored to
docs/PROJECT_DEEP_DIVE.md but general enough for similar docs.
"""

import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x4F, 0x46, 0xE5)
CODE_BG = "F2F2F5"
GREY = RGBColor(0x55, 0x55, 0x60)


def shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hexcolor)
    tc.append(shd)


def add_inline(paragraph, text):
    """Render **bold** and `code` spans inside a paragraph."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xC0, 0x30, 0x60)
        else:
            paragraph.add_run(part)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(10)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    # light grey shading via paragraph border/shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    pPr.append(shd)


def add_table(doc, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [
        [c.strip() for c in r.strip("|").split("|")]
        for r in rows[2:]  # skip the |---| separator row
    ]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ""
        add_inline(cell.paragraphs[0], h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r in body:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            if i < len(cells):
                cells[i].paragraphs[0].text = ""
                add_inline(cells[i].paragraphs[0], val)
    doc.add_paragraph()


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i]

        # fenced code block
        if line.strip().startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        # table
        if (
            line.strip().startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\s*\|[\s\-:|]+\|\s*$", lines[i + 1])
        ):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue

        # headings
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        # horizontal rule
        if line.strip() == "---":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "CCCCCC")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # blockquote
        if line.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i].lstrip("> ").rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_inline(p, " ".join(quote))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = GREY
            continue

        # bullet list
        if re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
            i += 1
            continue

        # numbered list
        if re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\s*\d+\. ", "", line))
            i += 1
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print(f"Wrote {docx_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
